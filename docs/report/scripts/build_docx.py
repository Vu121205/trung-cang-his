from pathlib import Path
import re
from docx import Document
from docx.shared import Cm, Pt
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.section import WD_SECTION

root = Path(__file__).resolve().parents[1]
source = root / 'bao-cao-thuc-tap.md'
output = root / 'TrungCang-HIS-bao-cao-thuc-tap.docx'
doc = Document()
section = doc.sections[0]
section.top_margin = Cm(2.5)
section.bottom_margin = Cm(2.5)
section.left_margin = Cm(4)
section.right_margin = Cm(2.5)
normal = doc.styles['Normal']
normal.font.name = 'Times New Roman'
normal.font.size = Pt(13)
normal.paragraph_format.line_spacing = 1.5
normal.paragraph_format.space_after = Pt(0)

def add_text(text, level=None):
    if level is None:
        p = doc.add_paragraph()
        p.paragraph_format.first_line_indent = Cm(1)
    else:
        p = doc.add_paragraph(style=f'Heading {min(level, 3)}')
        p.paragraph_format.first_line_indent = Cm(0)
    p.add_run(text)
    if level == 1:
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    return p

in_yaml = False
for raw in source.read_text(encoding='utf-8').splitlines():
    line = raw.strip()
    if line == '---':
        in_yaml = not in_yaml
        continue
    if in_yaml or not line:
        continue
    if line == r'\newpage':
        doc.add_page_break()
        continue
    image = re.match(r'!\[(.*?)\]\((.*?)\)', line)
    if image:
        caption, rel = image.groups()
        path = root / rel
        if path.suffix.lower() == '.svg':
            path = path.with_suffix('.png')
        if path.exists():
            p = doc.add_paragraph()
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER
            p.add_run().add_picture(str(path), width=Cm(15.5))
            cap = doc.add_paragraph(caption)
            cap.alignment = WD_ALIGN_PARAGRAPH.CENTER
            cap.runs[0].italic = True
        continue
    heading = re.match(r'^(#{1,3})\s+(.+)$', line)
    if heading:
        add_text(heading.group(2), len(heading.group(1)))
        continue
    if line.startswith('|'):
        continue
    if line.startswith('- '):
        p = doc.add_paragraph(style='List Bullet')
        p.add_run(line[2:])
        continue
    p = add_text(line)
    if line.startswith('*') and line.endswith('*'):
        p.runs[0].italic = True

doc.save(output)
print(output)
